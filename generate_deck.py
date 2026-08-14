from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Brand Colors ──────────────────────────────────────────────────
PRIMARY    = RGBColor(0x0A, 0x1F, 0x44)   # Deep navy
ACCENT     = RGBColor(0x00, 0x8B, 0xFF)   # Bright blue
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF4, 0xF6, 0xFA)
DARK_TEXT  = RGBColor(0x1A, 0x1A, 0x2E)
SUBTEXT    = RGBColor(0x55, 0x65, 0x80)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]   # completely blank


# ── Helper Functions ──────────────────────────────────────────────

def add_rect(slide, l, t, w, h, fill_color):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    return shape


def add_text(slide, text, l, t, w, h,
             font_size=18, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, italic=False):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox


def add_bullet_box(slide, items, l, t, w, h,
                   font_size=15, color=DARK_TEXT, bullet="▸  "):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = bullet + item
        run.font.size = Pt(font_size)
        run.font.color.rgb = color


def add_footer(slide):
    add_rect(slide, 0, 7.1, 13.33, 0.4, PRIMARY)
    add_text(slide, "www.adiakshsolutions.com", 0.3, 7.15, 6, 0.3,
             font_size=10, color=WHITE)


# ══════════════════════════════════════════════════════════════════
# SLIDE 1 — COVER
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, PRIMARY)
add_rect(slide, 0, 5.8, 13.33, 1.7, ACCENT)
add_rect(slide, 0.4, 0.35, 0.12, 6.8, ACCENT)

add_text(slide, "ADIAKSH SOLUTIONS",
         0.8, 1.5, 11, 1.2, font_size=44, bold=True, color=WHITE)
add_text(slide, "Your Digital & AI Engineering Partner",
         0.8, 2.75, 11, 0.7, font_size=22, color=ACCENT)
add_text(slide,
         "Web Development  •  Mobile Apps  •  AI Engineering  •  Cloud Infrastructure",
         0.8, 3.5, 11.5, 0.6, font_size=14, color=WHITE, italic=True)
add_text(slide, "www.adiakshsolutions.com",
         0.8, 6.0, 6, 0.5, font_size=13, color=WHITE)
add_text(slide, "Pune, India  |  Dubai, UAE",
         0.8, 6.55, 6, 0.5, font_size=13, color=WHITE)
add_text(slide, "Company Credentials & Capabilities",
         8.5, 6.55, 4.5, 0.5, font_size=12, color=DARK_TEXT,
         align=PP_ALIGN.RIGHT)

# ══════════════════════════════════════════════════════════════════
# SLIDE 2 — ABOUT US
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "About Adiaksh Solutions", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

add_text(slide,
         "Adiaksh Solutions is a boutique digital and technology studio headquartered "
         "in Pune, India, with a presence in Dubai, UAE. We partner with startups and "
         "global enterprises to design, build, and scale world-class digital products.",
         0.5, 1.3, 12.3, 1.2, font_size=15, color=DARK_TEXT)

stats = [
    ("Global Presence", "Pune & Dubai"),
    ("Fast Delivery",   "Agile & Iterative"),
    ("Quality First",   "Enterprise Grade"),
    ("Client Focused",  "Long-term Partners"),
]
for i, (title, sub) in enumerate(stats):
    x = 0.5 + i * 3.2
    add_rect(slide, x, 2.7, 2.9, 2.0, WHITE)
    add_text(slide, title, x + 0.1, 3.0,  2.7, 0.5, font_size=14, bold=True, color=PRIMARY)
    add_text(slide, sub,   x + 0.1, 3.55, 2.7, 0.45, font_size=12, color=SUBTEXT)

add_text(slide,
         "Our mission: Deliver measurable business outcomes through cutting-edge technology, "
         "superior design, and a relentless focus on client success.",
         0.5, 5.05, 12.3, 0.8, font_size=14, color=DARK_TEXT, italic=True)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 3 — OUR ROLE & RESPONSIBILITIES
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Our Role & Responsibilities", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

roles = [
    ("Strategic Technology Partner",
     ["Understand client business goals & translate them into tech solutions",
      "Advise on technology stack, architecture & scalability",
      "Provide roadmap planning & digital transformation guidance"]),
    ("End-to-End Product Development",
     ["Manage full software development lifecycle (SDLC)",
      "From concept & wireframes to development, QA & deployment",
      "Ensure on-time, on-budget delivery with agile sprints"]),
    ("Quality & Compliance Ownership",
     ["Maintain code quality, security best practices & performance",
      "Conduct regular code reviews, testing & vulnerability audits",
      "Adhere to industry standards (GDPR, OWASP, ISO best practices)"]),
    ("Ongoing Support & Maintenance",
     ["Post-launch monitoring, bug fixes & feature enhancements",
      "24/7 SLA-based support options available",
      "Dedicated account management for each client"]),
]

for i, (title, bullets) in enumerate(roles):
    row = i // 2
    col = i % 2
    x = 0.4 + col * 6.5
    y = 1.3 + row * 2.75
    add_rect(slide, x, y, 6.2, 2.4, WHITE)
    add_rect(slide, x, y, 6.2, 0.45, ACCENT)
    add_text(slide, title, x + 0.15, y + 0.05, 5.9, 0.38,
             font_size=13, bold=True, color=WHITE)
    add_bullet_box(slide, bullets, x + 0.15, y + 0.55, 5.9, 1.7,
                   font_size=12, color=DARK_TEXT)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 4 — CORE SERVICES
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Core Services & Offerings", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

services = [
    ("Web Development",
     "Full-stack web apps, SaaS platforms, e-commerce & enterprise portals "
     "using React, Next.js & cloud-native architectures."),
    ("Mobile App Development",
     "Cross-platform iOS & Android apps built with React Native / Flutter "
     "for seamless user experiences."),
    ("AI & Automation Engineering",
     "AI-powered products, ML pipelines, intelligent automation, chatbots "
     "& data-driven decision systems."),
    ("Cloud Infrastructure",
     "Architecture, deployment & management of scalable cloud solutions "
     "on AWS, Azure & GCP."),
    ("UX / UI Design",
     "User research, wireframing, prototyping & high-fidelity design "
     "that converts and delights."),
    ("Product Engineering",
     "Ideation to production: MVP development, product scaling, "
     "technical debt management & DevOps."),
]

for i, (title, desc) in enumerate(services):
    row = i // 3
    col = i % 3
    x = 0.4 + col * 4.3
    y = 1.25 + row * 2.85
    add_rect(slide, x, y, 4.0, 2.5, WHITE)
    add_rect(slide, x, y, 4.0, 0.08, ACCENT)
    add_text(slide, title, x + 0.15, y + 0.15, 3.7, 0.45,
             font_size=14, bold=True, color=PRIMARY)
    add_text(slide, desc,  x + 0.15, y + 0.65, 3.7, 1.7,
             font_size=11, color=DARK_TEXT)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 5 — TECHNICAL CAPABILITIES
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Technical Capabilities & Stack", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

tech = [
    ("Frontend",       ["React.js", "Next.js", "Vue.js", "TypeScript", "Tailwind CSS"]),
    ("Backend",        ["Node.js", "Python", "Django", "FastAPI", "REST & GraphQL APIs"]),
    ("Mobile",         ["React Native", "Flutter", "iOS (Swift)", "Android (Kotlin)"]),
    ("AI / ML",        ["OpenAI / GPT", "LangChain", "TensorFlow", "Scikit-learn", "RAG Pipelines"]),
    ("Cloud & DevOps", ["AWS", "Azure", "GCP", "Docker", "Kubernetes", "CI/CD"]),
    ("Databases",      ["PostgreSQL", "MongoDB", "MySQL", "Redis", "Firebase"]),
]

for i, (category, items) in enumerate(tech):
    row = i // 3
    col = i % 3
    x = 0.4 + col * 4.3
    y = 1.25 + row * 2.85
    add_rect(slide, x, y, 4.0, 2.5, WHITE)
    add_rect(slide, x, y, 4.0, 0.45, PRIMARY)
    add_text(slide, category, x + 0.15, y + 0.06, 3.7, 0.38,
             font_size=13, bold=True, color=WHITE)
    add_bullet_box(slide, items, x + 0.15, y + 0.55, 3.7, 1.8,
                   font_size=12, color=DARK_TEXT, bullet="• ")

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 6 — WHY CHOOSE US
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Why Partner with Adiaksh Solutions?", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

reasons = [
    ("Speed to Market",
     "Rapid prototyping & agile delivery cycles ensure your product reaches "
     "market faster than the competition."),
    ("Innovation-Driven",
     "We leverage the latest in AI, cloud & modern web technologies to keep "
     "your product ahead of the curve."),
    ("Cost-Effective",
     "Competitive pricing with Pune-based delivery center and Dubai-based "
     "client management — best of both worlds."),
    ("Deep Expertise",
     "Senior engineers and designers with 5-12 years of experience across "
     "diverse industries & technologies."),
    ("Global Standards",
     "We build to international quality, security, and compliance standards "
     "for clients across the globe."),
    ("True Partnership",
     "We don't just write code — we become a committed extension of your team, "
     "invested in your success."),
]

for i, (title, desc) in enumerate(reasons):
    row = i // 2
    col = i % 2
    x = 0.4 + col * 6.5
    y = 1.25 + row * 1.98
    add_rect(slide, x, y, 6.2, 1.75, WHITE)
    add_rect(slide, x, y, 0.08, 1.75, ACCENT)
    add_text(slide, title, x + 0.25, y + 0.15, 5.8, 0.42,
             font_size=14, bold=True, color=PRIMARY)
    add_text(slide, desc, x + 0.25, y + 0.6, 5.8, 0.9,
             font_size=12, color=DARK_TEXT)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 7 — OUR PROCESS
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Our Engagement & Delivery Process", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

steps = [
    ("01", "Discovery\n& Scoping",
     "Understand requirements, goals, timelines & budget"),
    ("02", "Architecture\n& Design",
     "UI/UX wireframes, system design & technology selection"),
    ("03", "Agile\nDevelopment",
     "Sprint-based development with regular demos & feedback"),
    ("04", "QA &\nTesting",
     "Automated & manual testing, performance & security review"),
    ("05", "Deployment\n& Go-Live",
     "CI/CD pipeline, cloud deployment & production launch"),
    ("06", "Support &\nGrowth",
     "Ongoing maintenance, feature roadmap & scaling support"),
]

for i, (num, title, desc) in enumerate(steps):
    x = 0.35 + i * 2.1
    add_rect(slide, x, 1.35, 1.85, 1.85, ACCENT)
    add_text(slide, num, x, 1.4, 1.85, 0.7,
             font_size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, title, x, 2.1, 1.85, 0.65,
             font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 5:
        add_text(slide, "->", x + 1.85, 2.0, 0.25, 0.5,
                 font_size=18, bold=True, color=PRIMARY)
    add_text(slide, desc, x, 3.35, 1.9, 1.0,
             font_size=11, color=DARK_TEXT, align=PP_ALIGN.CENTER)

add_text(slide,
         "Every project includes a dedicated project manager, weekly status reports, "
         "and full source-code ownership transferred to the client upon completion.",
         0.5, 5.0, 12.3, 0.8, font_size=13, color=PRIMARY, italic=True)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 8 — INDUSTRIES WE SERVE
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Industries We Serve", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

industries = [
    "E-Commerce & Retail",
    "Healthcare & MedTech",
    "FinTech & Banking",
    "EdTech & eLearning",
    "Enterprise SaaS",
    "Logistics & Supply Chain",
    "Travel & Hospitality",
    "Media & Entertainment",
]

for i, label in enumerate(industries):
    row = i // 4
    col = i % 4
    x = 0.5 + col * 3.1
    y = 1.3 + row * 2.7
    add_rect(slide, x, y, 2.75, 2.2, WHITE)
    add_rect(slide, x, y, 2.75, 0.08, ACCENT)
    add_text(slide, label, x + 0.1, y + 0.5, 2.55, 1.5,
             font_size=13, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 9 — ENGAGEMENT MODELS
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.1, PRIMARY)
add_text(slide, "Engagement Models", 0.5, 0.2, 12, 0.75,
         font_size=26, bold=True, color=WHITE)

models = [
    ("Fixed Price Project",
     ["Well-defined scope & deliverables",
      "Fixed timeline & cost",
      "Ideal for MVPs and defined projects",
      "Regular milestone-based reviews"]),
    ("Time & Material (T&M)",
     ["Flexible, evolving requirements",
      "Pay only for hours worked",
      "Full transparency with timesheets",
      "Best for long-term development"]),
    ("Dedicated Team",
     ["Dedicated developers & designers",
      "Works as an extension of your team",
      "Full control over priorities",
      "Monthly retainer-based pricing"]),
]

for i, (title, bullets) in enumerate(models):
    x = 0.5 + i * 4.2
    add_rect(slide, x, 1.3, 3.9, 5.0, WHITE)
    add_rect(slide, x, 1.3, 3.9, 0.08, ACCENT)
    add_text(slide, title, x, 1.5, 3.9, 0.55,
             font_size=15, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_bullet_box(slide, bullets, x + 0.2, 2.2, 3.5, 3.0,
                   font_size=13, color=DARK_TEXT)

add_footer(slide)

# ══════════════════════════════════════════════════════════════════
# SLIDE 10 — CALL TO ACTION / CONTACT
# ══════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
add_rect(slide, 0, 0, 13.33, 7.5, PRIMARY)
add_rect(slide, 0, 5.5, 13.33, 2.0, ACCENT)
add_rect(slide, 0.4, 0.35, 0.12, 6.8, ACCENT)

add_text(slide, "Let's Build Something Great Together",
         0.8, 1.2, 11.5, 1.2, font_size=36, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER)
add_text(slide,
         "We're ready to understand your challenges and propose the right digital solution.\n"
         "Schedule a free consultation call with our team today.",
         0.8, 2.6, 11.5, 1.0, font_size=16, color=ACCENT,
         align=PP_ALIGN.CENTER, italic=True)

add_text(slide, "www.adiakshsolutions.com",
         1.5, 3.85, 5.0, 0.55, font_size=15, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, "Pune, India  |  Dubai, UAE",
         7.0, 3.85, 5.0, 0.55, font_size=15, color=WHITE, align=PP_ALIGN.CENTER)

add_text(slide, "2025 Adiaksh Solutions. All rights reserved.",
         0.5, 6.85, 12.3, 0.4, font_size=11, color=DARK_TEXT,
         align=PP_ALIGN.CENTER)

# ── Save ──────────────────────────────────────────────────────────
output_path = "Adiaksh_Solutions_Company_Deck.pptx"
prs.save(output_path)
print(f"Presentation saved -> {output_path}")
